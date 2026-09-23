# -*- coding: utf-8 -*-
"""
CGA_EmptyMyRoom_KO_v1.md  ->  IEEE CG&A 매거진 템플릿(.docx) 변환기.

원본 템플릿(CGA_CS_Mag_2_11_26.docx)의 스타일·섹션(2단 조판)·머리글/바닥글을 그대로
유지한 채, 제목/저자/초록/본문/참고문헌/저자약력만 교체한다.

사용법:  python build_docx.py <template.docx> <input.md> <output.docx>
"""
import re
import sys
import docx
from docx.oxml.ns import qn
from docx.shared import Pt

EAST_ASIA_FONT = '맑은 고딕'


def _set_korean_font(run):
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = rPr.makeelement(qn('w:rFonts'), {})
        rPr.insert(0, rFonts)
    rFonts.set(qn('w:eastAsia'), EAST_ASIA_FONT)


def add_runs(par, text):
    """**bold** 마크업만 해석해서 run으로 나눠 넣는다."""
    for i, chunk in enumerate(re.split(r'(\*\*[^*]+\*\*)', text)):
        if not chunk:
            continue
        if chunk.startswith('**') and chunk.endswith('**'):
            run = par.add_run(chunk[2:-2])
            run.bold = True
        else:
            run = par.add_run(chunk)
        _set_korean_font(run)
    return par


def parse_md(path):
    """마크다운을 (섹션명 -> 줄 리스트) 딕셔너리로 읽는다."""
    sections, cur = {}, None
    for line in open(path, encoding='utf-8'):
        line = line.rstrip('\n')
        if line.startswith('<!--'):
            continue
        m = re.match(r'^# ([A-Z\-]+)$', line)
        if m:
            cur = m.group(1)
            sections[cur] = []
            continue
        if cur:
            sections[cur].append(line)
    return sections


def all_paragraphs(doc):
    """표·텍스트박스 안까지 포함한 모든 문단(XML 순회)."""
    from docx.text.paragraph import Paragraph
    return [Paragraph(el, doc) for el in doc.element.body.iter(qn('w:p'))]


def par_style_id(el):
    st = el.find(qn('w:pPr') + '/' + qn('w:pStyle'))
    return st.get(qn('w:val')) if st is not None else None


def clear_par(par):
    """하이퍼링크·북마크를 포함한 문단 내용 전체 제거."""
    for child in list(par._element):
        if child.tag != qn('w:pPr'):
            par._element.remove(child)


def replace_title_block(doc, article_type, title, authors):
    pars = all_paragraphs(doc)

    for par in pars:
        if par_style_id(par._element) == 'Titledocument':
            clear_par(par)
            add_runs(par, title)
            break

    author_pars = [p for p in pars if par_style_id(p._element) == 'Authors' and p.text.strip()]
    for i, par in enumerate(author_pars):
        clear_par(par)
        if i < len(authors):
            add_runs(par, authors[i])
    # 템플릿(3명)보다 저자가 많으면 마지막 문단을 복제해 이어 붙인다
    if author_pars and len(authors) > len(author_pars):
        import copy
        anchor = author_pars[-1]._element
        for extra in authors[len(author_pars):]:
            new_el = copy.deepcopy(anchor)
            anchor.addnext(new_el)
            anchor = new_el
            from docx.text.paragraph import Paragraph
            par = Paragraph(new_el, doc)
            clear_par(par)
            add_runs(par, extra)

    done = False
    for par in all_paragraphs(doc):
        if 'Article Type: Description' in par.text:
            clear_par(par)
            if not done:
                add_runs(par, article_type)
                done = True


def clear_body(doc):
    """초록 뒤 sectPr 문단까지 남기고, 최종 sectPr 앞의 본문을 전부 제거."""
    body = doc.element.body
    children = list(body)
    # 초록 영역을 닫는 두 번째 sectPr 보유 문단의 인덱스를 찾는다
    sect_idx = [i for i, c in enumerate(children)
                if c.tag == qn('w:p') and c.find(qn('w:pPr') + '/' + qn('w:sectPr')) is not None]
    keep_until = sect_idx[1]        # 0: 제목블록, 1: 초록블록
    for c in children[keep_until + 1:]:
        if c.tag == qn('w:sectPr'):  # 최종 sectPr(2단 조판)는 보존
            continue
        body.remove(c)


def replace_abstract(doc, text):
    for par in all_paragraphs(doc):
        if par_style_id(par._element) == 'Abstract':
            clear_par(par)
            add_runs(par, text)
            return


def add_par(doc, style, text):
    par = doc.add_paragraph(style=style)
    add_runs(par, text)
    return par


def add_table(doc, rows):
    tbl = doc.add_table(rows=len(rows), cols=len(rows[0]))
    tbl.style = doc.styles['Table Grid']
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.text = ''
            par = cell.paragraphs[0]
            try:
                par.style = doc.styles['Table Paragraph']
            except KeyError:
                pass
            run = par.add_run(val)
            run.font.size = Pt(8)
            if r == 0:
                run.bold = True
            _set_korean_font(run)
    return tbl


def render_body(doc, lines):
    prev_is_heading = True
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue

        # 마크다운 표
        if line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].startswith('|'):
                cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-{2,}:?', c) for c in cells):
                    rows.append(cells)
                i += 1
            add_table(doc, rows)
            prev_is_heading = False
            continue

        if line.startswith('### '):
            add_par(doc, 'Heading 2', line[4:].strip())
            prev_is_heading = True
        elif line.startswith('## '):
            add_par(doc, 'Heading 1', line[3:].strip().upper() if line[3:].strip().isascii()
                    else line[3:].strip())
            prev_is_heading = True
        elif re.match(r'^\*\*FIGURE \d+\.\*\*', line):
            add_par(doc, 'FigureCaption', line.replace('**', ''))
            prev_is_heading = True
        elif re.match(r'^\*\*TABLE \d+\.\*\*', line):
            add_par(doc, 'TableCaption', line.replace('**', ''))
            prev_is_heading = True
        elif line.startswith('* '):
            add_par(doc, 'List Paragraph', line[2:].strip())
            prev_is_heading = False
        elif line.startswith('    ') and line.strip():
            add_par(doc, 'DisplayFormula', line.strip())
            prev_is_heading = False
        else:
            add_par(doc, 'Para' if prev_is_heading else 'ParaContinue', line.strip())
            prev_is_heading = False
        i += 1


def main():
    template, md_path, out_path = sys.argv[1], sys.argv[2], sys.argv[3]
    s = parse_md(md_path)
    txt = lambda key: '\n'.join(s.get(key, [])).strip()
    lines = lambda key: [l for l in s.get(key, []) if l.strip()]

    doc = docx.Document(template)
    replace_title_block(doc, txt('ARTICLE-TYPE'), txt('TITLE'), lines('AUTHORS'))
    replace_abstract(doc, txt('ABSTRACT'))
    clear_body(doc)

    render_body(doc, s.get('BODY', []))

    add_par(doc, 'AckHead', 'ACKNOWLEDGMENTS')
    for l in lines('ACKNOWLEDGMENTS'):
        add_par(doc, 'Para', l)

    add_par(doc, 'ReferenceHead', 'REFERENCES')
    for l in lines('REFERENCES'):
        add_par(doc, 'Bib_entry', l)

    for l in lines('AUTHOR-BIOS'):
        add_par(doc, 'AuthorBio', l)

    doc.save(out_path)
    print(f"saved: {out_path}")


if __name__ == '__main__':
    main()
