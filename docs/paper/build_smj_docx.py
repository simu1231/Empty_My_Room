# -*- coding: utf-8 -*-
"""
SMJ_EmptyMyRoom_KO_v1.md -> 스마트미디어저널(영문) 논문서식 .docx 변환기.

원본 서식(스마트미디어저널 논문서식 영문.doc를 .docx로 변환한 것)의 제목표·각주표·
저자약력표·2단 섹션 구성을 그대로 유지한 채 내용만 교체한다.

사용법: python build_smj_docx.py <template.docx> <input.md> <output.docx>
"""
import re
import sys
import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH as AL
from docx.oxml.ns import qn
from docx.shared import Pt

# (스타일, 정렬, 글꼴, 크기pt, 굵게)
FMT = {
    'head1':   ('s0',     AL.CENTER,  'HY중고딕',   13,  True),
    'head2':   ('s0',     AL.LEFT,    'HY중고딕',   11,  True),
    'body':    ('s0',     AL.JUSTIFY, 'HY신명조',   11,  False),
    'figcap':  ('s0',     AL.CENTER,  'HY중고딕',   9.5, False),
    'tabcap':  ('그림/표', AL.CENTER,  'HY중고딕',   9.5, False),
    'refhead': ('Normal', AL.CENTER,  '한양중고딕', 11,  True),
    'ref':     ('Normal', AL.LEFT,    'HY신명조',   9.5, False),
    'bio':     ('Normal', AL.LEFT,    'HY신명조',   9,   False),
}


def style_run(run, font, size, bold):
    run.font.size = Pt(size)
    run.bold = bold
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = rPr.makeelement(qn('w:rFonts'), {})
        rPr.insert(0, rFonts)
    for attr in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        rFonts.set(qn(attr), font)


def fill_par(par, text, kind, style=True):
    s, align, font, size, bold = FMT[kind]
    if style:
        try:
            par.style = par.part.document.styles[s]
        except KeyError:
            pass
    par.alignment = align
    for chunk in re.split(r'(\*\*[^*]+\*\*)', text):
        if not chunk:
            continue
        b = bold
        if chunk.startswith('**') and chunk.endswith('**'):
            chunk, b = chunk[2:-2], True
        style_run(par.add_run(chunk), font, size, b)
    return par


def clear_par(par):
    for child in list(par._element):
        if child.tag != qn('w:pPr'):
            par._element.remove(child)


def insert_par(doc, anchor, text, kind):
    par = doc.add_paragraph()
    anchor.addprevious(par._element)
    return fill_par(par, text, kind)


def insert_table(doc, anchor, rows):
    tbl = doc.add_table(rows=len(rows), cols=len(rows[0]))
    try:
        tbl.style = doc.styles['Table Grid']
    except KeyError:
        pass
    anchor.addprevious(tbl._tbl)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            par = tbl.cell(r, c).paragraphs[0]
            par.alignment = AL.CENTER
            style_run(par.add_run(val), 'HY신명조', 8, r == 0)
    return tbl


def parse_md(path):
    sections, cur = {}, None
    for line in open(path, encoding='utf-8'):
        line = line.rstrip('\n')
        if line.startswith('<!--'):
            continue
        m = re.match(r'^# ([A-Z\-]+)$', line)
        if m:
            cur = m.group(1)
            sections[cur] = []
            continue
        if cur is not None:
            sections[cur].append(line)
    return sections


def main():
    template, md_path, out_path = sys.argv[1], sys.argv[2], sys.argv[3]
    s = parse_md(md_path)
    txt = lambda k: '\n'.join(s.get(k, [])).strip()
    rows = lambda k: [l.strip() for l in s.get(k, []) if l.strip()]

    doc = docx.Document(template)
    t_title, t_note, t_bio = doc.tables[0], doc.tables[1], doc.tables[3]

    # ── 제목표: 제목 / 저자 / 초록+키워드 ──────────────────────────────
    def set_cell(cell, items):
        """items = [(text, kind)] — 첫 문단을 재사용하고 나머지는 추가."""
        for p in cell.paragraphs[1:]:
            p._element.getparent().remove(p._element)
        base = cell.paragraphs[0]
        clear_par(base)
        fill_par(base, items[0][0], items[0][1], style=False)
        for text, kind in items[1:]:
            fill_par(cell.add_paragraph(), text, kind, style=False)

    FMT['title'] = ('s0', AL.CENTER, 'HY신명조', 17, False)
    FMT['titleen'] = ('s0', AL.CENTER, 'HY신명조', 12, False)
    FMT['authors'] = ('s0', AL.CENTER, 'HY신명조', 11, True)
    FMT['abs'] = ('s0', AL.JUSTIFY, 'HY신명조', 10, False)
    FMT['kw'] = ('s0', AL.LEFT, 'HY신명조', 10, False)
    FMT['note'] = ('s0', AL.LEFT, 'HY중고딕', 8.5, False)

    set_cell(t_title.cell(2, 1), [(txt('TITLE'), 'title'), (txt('TITLE-EN'), 'titleen')])
    set_cell(t_title.cell(4, 1), [(txt('AUTHORS') + '  (' + txt('AUTHORS-EN') + ')', 'authors')])
    set_cell(t_title.cell(7, 1), [
        (txt('ABSTRACT'), 'abs'),
        ('', 'abs'),
        ('Keywords : ' + txt('KEYWORDS'), 'kw'),
        ('Keywords : ' + txt('KEYWORDS-EN'), 'kw'),
    ])

    # ── 각주표: 지원사업 / 투고일 / 교신저자 ───────────────────────────
    note = rows('FOOTNOTE')
    split = next((i for i, l in enumerate(note) if l.startswith('Manuscript')), len(note))
    set_cell(t_note.cell(0, 0), [(l, 'note') for l in note[:split]] or [('', 'note')])
    set_cell(t_note.cell(1, 0), [(l, 'note') for l in note[split:]] or [('', 'note')])

    # ── 기존 본문 제거: 각주표와 저자약력표 사이를 비운다 ──────────────
    body = doc.element.body
    kids = list(body)
    i0, i1 = kids.index(t_note._tbl), kids.index(t_bio._tbl)
    for child in kids[i0 + 1:i1]:
        body.remove(child)
    # 제목 블록을 닫는 sectPr 문단과 각주표(부동 배치) 사이에 남은 템플릿 견본 문단
    # ('INTRODUCTION' 제목 등)도 제거한다. 제거하지 않으면 마크다운이 넣는 제목과
    # 중복된다. 각주표는 tblpPr로 고정 배치되므로 흐름상 위치를 옮겨도 무방하다.
    sect_idx = next(i for i, c in enumerate(kids)
                    if c.tag == qn('w:p')
                    and c.find(qn('w:pPr') + '/' + qn('w:sectPr')) is not None)
    for child in kids[sect_idx + 1:i0]:
        body.remove(child)

    # ── 본문 삽입 ──────────────────────────────────────────────────────
    anchor = t_bio._tbl
    lines = s.get('BODY', [])
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        if line.startswith('|'):
            table = []
            while i < len(lines) and lines[i].startswith('|'):
                cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-{2,}:?', c) for c in cells):
                    table.append(cells)
                i += 1
            insert_table(doc, anchor, table)
            continue
        if line.startswith('### '):
            insert_par(doc, anchor, line[4:].strip(), 'head2')
        elif line.startswith('## '):
            insert_par(doc, anchor, line[3:].strip(), 'head1')
        elif re.match(r'^\*\*Fig\. \d+\.\*\*', line):
            insert_par(doc, anchor, line.replace('**', ''), 'figcap')
        elif re.match(r'^\*\*Table \d+\.\*\*', line):
            insert_par(doc, anchor, line.replace('**', ''), 'tabcap')
        elif line.startswith('* '):
            insert_par(doc, anchor, '· ' + line[2:].strip(), 'body')
        else:
            insert_par(doc, anchor, line.strip(), 'body')
        i += 1

    insert_par(doc, anchor, 'REFERENCES', 'refhead')
    for n, ref in enumerate(rows('REFERENCES'), 1):
        insert_par(doc, anchor, f'[{n}] {ref}', 'ref')

    # ── 저자약력표 ─────────────────────────────────────────────────────
    bios = rows('AUTHOR-BIOS')
    items = []
    for line in bios:
        items.append((line, 'bio'))
    set_cell(t_bio.cell(3, 0), items or [('', 'bio')])

    doc.save(out_path)
    print(f"saved: {out_path}")


if __name__ == '__main__':
    main()
